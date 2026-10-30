# Builds LocalBoost-Pitch.pptx from the hackathon template + app screenshots (docs/deck).
#   python scripts/build-deck.py
import os, copy
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
TEMPLATE = r'C:\Users\Parlit___\Documents\Template-Pitch-Hackathon.pptx'
SHOTS = os.path.join(ROOT, 'docs', 'deck')
CROPS = os.path.join(SHOTS, 'crops'); os.makedirs(CROPS, exist_ok=True)
OUT = os.path.join(ROOT, 'LocalBoost-Pitch.pptx')
LOGO = os.path.join(ROOT, 'dist', 'icons', 'logo.png')

BG, PANEL, ACC, GOLD, GREEN = '192B37', '0F1F29', 'FF5640', 'FFD700', '2ECC71'
WHITE, MUTED, LINE = 'FFFFFF', 'A9B7C2', '2C4152'
SANS, MONO = 'IBM Plex Sans', 'Fira Code'
rgb = lambda h: RGBColor.from_string(h)

prs = Presentation(TEMPLATE)
# drop every template slide
lst = prs.slides._sldIdLst
for sid in list(lst):
    prs.part.drop_rel(sid.rId); lst.remove(sid)
BLANK = next(l for l in prs.slide_layouts if l.name == 'Blank')


def crop(name, box, out=None):
    im = Image.open(os.path.join(SHOTS, name + '.png')).convert('RGB').crop(box)
    p = os.path.join(CROPS, (out or name) + '.png'); im.save(p); return p, im.size


def new_slide(notes=''):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    if notes: s.notes_slide.notes_text_frame.text = notes
    return s


def rect(s, x, y, w, h, fill=PANEL, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill: r.fill.solid(); r.fill.fore_color.rgb = rgb(fill)
    else: r.fill.background()
    if line: r.line.color.rgb = rgb(line); r.line.width = Pt(0.75)
    else: r.line.fill.background()
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE: r.adjustments[0] = radius
    r.shadow.inherit = False
    return r


def text(s, x, y, w, h, t, size=14, color=WHITE, bold=False, font=SANS, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.1):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(t if isinstance(t, list) else [t]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        r = p.add_run(); r.text = line; f = r.font
        f.size = Pt(size); f.bold = bold; f.name = font; f.color.rgb = rgb(color)
    return tb


def head(s, kicker, title, n):
    text(s, 0.6, 0.45, 9, 0.25, kicker.upper(), 11, ACC, font=MONO)
    text(s, 0.6, 0.78, 12.1, 0.7, title, 30, WHITE, True)
    text(s, 11.6, 0.45, 1.15, 0.25, f'{n:02d} / 14', 11, MUTED, font=MONO, align=PP_ALIGN.RIGHT)


def card(s, x, y, w, h, title=None, body=None, accent=ACC, size=13):
    rect(s, x, y, w, h); rect(s, x, y, 0.04, h, accent)
    cy = y + 0.2
    if title:
        text(s, x + 0.25, cy, w - 0.45, 0.3, title, 15, WHITE, True); cy += 0.42
    if body:
        text(s, x + 0.25, cy, w - 0.45, h - (cy - y) - 0.1, body, size, MUTED, spacing=1.15)


def pic(s, path_size, x, y, w=None, h=None, frame=True):
    path, (iw, ih) = path_size
    if w and not h: h = w * ih / iw
    if h and not w: w = h * iw / ih
    if frame: rect(s, x - 0.04, y - 0.04, w + 0.08, h + 0.08, fill=LINE)
    s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    return w, h


def phone(s, name, x, y, h, box=None):
    """Phone screenshot in a rounded device frame, height h."""
    p, (iw, ih) = crop(name, box or (0, 0, 780, 1688), 'ph-' + name)
    w = h * iw / ih
    f = rect(s, x - 0.1, y - 0.1, w + 0.2, h + 0.2, fill='0A141B', line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
    s.shapes.add_picture(p, Inches(x), Inches(y), Inches(w), Inches(h))
    return w


def label(s, x, y, t, color=MUTED):
    text(s, x, y, 5, 0.25, t, 10, color, font=MONO)


def bar(s, t, n=None):
    rect(s, 0.6, 6.85, 12.13, 0.38, fill=PANEL)
    text(s, 0.8, 6.85, 11.7, 0.38, t, 12, MUTED, anchor=MSO_ANCHOR.MIDDLE)


# ---------------------------------------------------------------- 1 title
s = new_slide('TU: bună ziua. Noi suntem echipa din spatele LocalBoost. În trei minute vă arătăm cum o firmă mică devine recomandată online, fără buget mare de marketing.')
s.shapes.add_picture(LOGO, Inches(5.92), Inches(1.2), Inches(1.5), Inches(1.5))
text(s, 0, 2.95, 13.333, 0.3, 'HACKATHON · PITCH', 12, ACC, font=MONO, align=PP_ALIGN.CENTER)
text(s, 0, 3.3, 13.333, 0.9, 'LocalBoost', 54, WHITE, True, align=PP_ALIGN.CENTER)
text(s, 0, 4.25, 13.333, 0.4, 'Signal-to-Action OS pentru afacerile locale', 20, GOLD, align=PP_ALIGN.CENTER)
rect(s, 6.33, 4.9, 0.67, 0.03, ACC)
text(s, 2.5, 5.15, 8.33, 0.8, 'Găsim golurile din prezența online a unei firme și le reparăm, ca să fie recomandată online.', 16, MUTED, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- 2 clients
s = new_slide('TU: ne adresăm managerilor de firme, fondatorilor și antreprenorilor. Oameni care știu să-și facă treaba, dar nu au timp sau buget pentru marketing.')
head(s, 'Clienții · 1', 'Cui ne adresăm', 2)
for i, (t, b) in enumerate([('Manageri de firme', 'Conduc o afacere locală și răspund de fiecare lună: clienți, cheltuieli, reputație.'),
                            ('Fondatori', 'Au pornit o afacere nouă și trebuie să fie găsiți de oameni care nu-i cunosc încă.'),
                            ('Antreprenori', 'Fac totul singuri și nu au timp de analize, rapoarte sau agenții.')]):
    card(s, 0.6 + i * 4.1, 1.9, 3.95, 2.1, t, b)
card(s, 0.6, 4.3, 12.13, 1.9, 'Ce au în comun', ['Nu au un buget mare de marketing.', 'Au deja recenzii, postări și clienți, dar nimeni nu le adună într-un singur loc.', 'Au nevoie de un pas următor clar, nu de încă un grafic.'], GOLD)
bar(s, 'Demo-ul din această prezentare: Cafeneaua Copou, Iași — o afacere fictivă.')

# ---------------------------------------------------------------- 3 problem
s = new_slide('TU: problema. Firmele mici depind de bugete mari de marketing ca să fie văzute. Fără buget, informația despre ele e împrăștiată și ele nu apar când cineva caută.')
head(s, 'Problema · 2', 'Vizibilitatea costă, iar firmele mici nu au bugetul', 3)
pts = [('Depind de buget', 'Cine nu poate plăti reclame și agenții rămâne greu de găsit online.'),
       ('Informație împrăștiată', 'Recenzii, comentarii, rețele sociale și site sunt în locuri diferite.'),
       ('Goluri nevăzute', 'Un site fără buton de acțiune sau un program neclar pierd clienți fără ca managerul să afle.')]
for i, (t, b) in enumerate(pts):
    card(s, 0.6, 1.8 + i * 1.5, 6.6, 1.35, t, b)
ph = phone(s, 'p-signals', 8.2, 1.55, 4.9, (0, 0, 780, 1688))
phone(s, 'p-competitors', 10.9, 1.55, 4.9) if False else None
text(s, 0.6, 6.4, 7, 0.3, 'Ecran real din aplicație: semnalele unei firme, adunate într-un loc.', 11, MUTED)

# ---------------------------------------------------------------- 4 solution
s = new_slide('TU: soluția. O platformă care găsește golurile din prezența online a firmei și le repară, ca firma să poată fi recomandată online. Gimi vă arată cum.')
head(s, 'Soluția · 3', 'Observă → Diagnostichează → Acționează → Măsoară', 4)
steps = [('1', 'Observă', 'Adună semnale: recenzii, comentarii de pe social media, site, căutări.'),
         ('2', 'Diagnostichează', 'Găsește golurile și alege o oportunitate cu dovezi.'),
         ('3', 'Acționează', 'Creează oferta, postarea, titlurile și un site temporar.'),
         ('4', 'Măsoară', 'Compară înainte și după, la 48 de ore, cu cifre etichetate.')]
for i, (n, t, b) in enumerate(steps):
    x = 0.6 + i * 3.08
    rect(s, x, 1.9, 2.93, 2.6); rect(s, x, 1.9, 2.93, 0.04, ACC if i < 3 else GREEN)
    text(s, x + 0.25, 2.1, 1, 0.3, f'0{n}', 12, ACC, font=MONO)
    text(s, x + 0.25, 2.45, 2.5, 0.4, t, 18, WHITE, True)
    text(s, x + 0.25, 3.0, 2.45, 1.4, b, 13, MUTED, spacing=1.15)
p = crop('d-overview', (260, 60, 1440, 700), 'overview-top')
pic(s, p, 0.6, 4.75, h=1.95)
bar(s, 'Fiecare cifră din aplicație are sursa afișată la hover sau long press.')

# ---------------------------------------------------------------- 5 observe
s = new_slide('GIMI: pasul unu, Observă. Managerul scrie numele firmei, orașul și linkurile: site, Instagram, TikTok și Facebook. Dacă nu are una dintre ele, apasă Nu am și aplicația îi propune idei de nume.')
head(s, 'Funcții · Observă', 'Începem de la ce știe managerul', 5)
pic(s, crop('d-form', (260, 60, 1440, 1240), 'form'), 0.6, 1.7, h=4.95)
phone(s, 'p-overview', 10.5, 1.65, 4.7) if False else None
card(s, 5.8, 1.7, 3.85, 2.4, 'Date introduse de manager', ['Nume, tip de afacere, oraș', 'Website, Instagram, TikTok, Facebook', 'Buton „Nu am” pentru fiecare'], ACC)
phone(s, 'p-overview', 10.3, 1.6, 4.9)
card(s, 5.8, 4.3, 3.85, 2.35, 'Facebook adăugat', 'Pagina de Facebook intră în analiză la fel ca Instagram și TikTok, inclusiv în footerul site-ului generat.', GOLD)
bar(s, 'Tot ce lipsește devine o sarcină: aplicația nu inventează date.')

# ---------------------------------------------------------------- 6 signals
s = new_slide('GIMI: semnalele. Pe lângă recenzii, aplicația adună și comentarii de pe Instagram, TikTok și Facebook. Atenție: în demo sunt comentarii-exemplu, etichetate. Pentru comentarii reale e nevoie de conectarea oficială a conturilor, aprobată de proprietar.')
head(s, 'Funcții · Semnale', 'Recenzii și comentarii din social media', 6)
pic(s, crop('d-signals', (260, 60, 1440, 1160), 'signals'), 0.6, 1.7, h=4.95)
phone(s, 'p-signals', 10.3, 1.6, 4.9) if False else None
card(s, 7.1, 1.7, 5.6, 1.5, 'Surse de semnale', 'Recenzii, întrebări ale clienților, audit de site, tendințe de căutare și comentarii sociale.', ACC, 13)
card(s, 7.1, 3.4, 5.6, 1.65, 'Spunem deschis ce e simulat', 'Comentariile din demo sunt exemple marcate DEMO. Datele reale vin doar prin conexiuni oficiale sau import CSV.', GOLD, 13)
card(s, 7.1, 5.25, 5.6, 1.4, 'Platformele lipsă sunt listate', 'Dacă firma nu are TikTok, aplicația spune că niciun cont nu a fost adăugat.', GREEN, 13)

# ---------------------------------------------------------------- 7 diagnosis
s = new_slide('GIMI: diagnosticul. Din semnale aplicația scoate ce apreciază clienții, ce lipsește și care este cea mai bună oportunitate. Aici, un coffee break între 14:00 și 18:00.')
head(s, 'Funcții · Diagnostic', 'Ce spun clienții și unde se pierd ei', 7)
pic(s, crop('d-overview', (260, 60, 1440, 1700), 'overview'), 0.6, 1.7, h=4.95)
phone(s, 'p-pulse', 9.0, 1.6, 4.9)
card(s, 5.15, 1.7, 3.2, 2.2, 'Review pulse', 'Teme din recenzii, cu procente și motive.', ACC, 12)
card(s, 5.15, 4.1, 3.2, 2.55, 'Reach dinamic', 'Alegi intervalul (7, 30, 90 de zile sau data dorită) și cifrele se recalculează.', GOLD, 12)
bar(s, 'AI-ul poate greși: avertismentul apare în aplicație înainte de orice acțiune.') if False else None

# ---------------------------------------------------------------- 8 competitors
s = new_slide('GIMI: concurența. O pagină separată compară firma cu cinci locuri din zonă: rating, recenzii, reach, rezervări online și ofertă de după-amiază. Locurile din demo sunt fictive.')
head(s, 'Funcții · Concurența', 'Unde stai față de vecini', 8)
pic(s, crop('d-competitors', (260, 60, 1440, 1560), 'comp1'), 0.6, 1.7, h=4.95)
phone(s, 'p-competitors', 9.4, 1.6, 4.9)
card(s, 5.2, 1.7, 3.7, 2.45, 'Comparație', 'Bare și grafic radar: tu față de piață.', ACC, 12)
card(s, 5.2, 4.35, 3.7, 2.3, 'Idei', 'Sfaturi concrete pentru golurile găsite.', GOLD, 12)

# ---------------------------------------------------------------- 9 act
s = new_slide('GIMI: acționează. Mission Control pornește misiunea de 48 de ore: cinci agenți AI lucrează pe rând. Dacă AI-ul nu răspunde în 8 secunde, intră un răspuns local, iar aplicația spune clar asta.')
head(s, 'Funcții · Acționează', 'Mission Control: o misiune de 48 de ore', 9)
pic(s, crop('d-mission', (260, 60, 1440, 1240), 'mission'), 0.6, 1.7, h=4.95)
phone(s, 'p-mission', 9.9, 1.6, 4.9) if False else None
card(s, 6.7, 1.7, 2.9, 2.45, '5 agenți', 'Analist, strateg, creator, builder și Name Scout.', ACC, 12)
card(s, 6.7, 4.35, 2.9, 2.3, 'Fallback local', 'Fără răspuns AI în 8 s, aplicația continuă cu reguli locale.', GOLD, 12)
phone(s, 'p-mission', 10.4, 1.6, 4.9)

# ---------------------------------------------------------------- 10 website
s = new_slide('GIMI: site-ul temporar. Pentru firmele fără site, generăm o pagină cu ofertă și buton de rezervare. Managerul editează textul, culorile, fontul și o exportă.')
head(s, 'Funcții · Acționează', 'Un site temporar, gata în câteva secunde', 10)
pic(s, crop('d-website', (260, 60, 1440, 1290), 'website'), 0.6, 1.7, h=4.95)
phone(s, 'p-website', 9.9, 1.6, 4.9)
card(s, 6.2, 1.7, 3.1, 2.1, 'Editabil', 'Titlu, text, culori, temă și font, cu previzualizare pe desktop, tabletă și telefon.', ACC, 12)
card(s, 6.2, 4.0, 3.1, 2.65, 'Export', 'ZIP, HTML sau PNG. Linkurile sociale, inclusiv Facebook, apar în footer.', GOLD, 12)

# ---------------------------------------------------------------- 11 pack + proof
s = new_slide('GIMI: măsoară. Pachetul de acțiune are oferta, postarea, titlurile și merch-ul, toate cu un buton clar. Proof Board compară înainte cu după, iar valorile sunt introduse de manager și marcate DEMO.')
head(s, 'Funcții · Măsoară', 'Pachet de acțiune și Proof Board', 11)
pic(s, crop('d-pack', (260, 60, 1440, 1450), 'pack'), 0.6, 1.7, h=4.95)
pic(s, crop('d-proof', (260, 60, 1440, 1000), 'proof'), 5.2, 1.7, w=4.4)
phone(s, 'p-pulse', 10.4, 1.6, 4.9)
card(s, 5.2, 5.4, 4.4, 1.25, 'Fără promisiuni goale', 'Ținte, nu predicții. Valorile după 48 de ore sunt introduse de manager și etichetate.', GREEN, 12)

# ---------------------------------------------------------------- 12 transformation
s = new_slide('TU: transformarea. Managerul primește verificarea recenziilor și sfaturi, o sinteză a tot ce se spune online despre firma lui. Rezultatul: mai multă încredere, dacă informațiile sunt verificate, mai mulți clienți și o economie locală mai puternică.')
head(s, 'Transformarea · 4', 'Din semnale împrăștiate în clienți', 12)
items = [('Verificare și sfaturi', 'Recenziile sunt analizate, iar managerul primește pași concreți.', ACC),
         ('O sinteză a firmei', 'Tot ce există online despre afacere, într-un singur loc.', ACC),
         ('Mai multă încredere', 'Clienții au încredere în firmă când informațiile sunt verificate.', GOLD),
         ('Mai mulți clienți', 'Firma devine ușor de găsit și de recomandat online.', GOLD),
         ('Economie locală mai puternică', 'Startup-uri și afaceri mici, mai vizibile în orașul lor.', GREEN)]
for i, (t, b, c) in enumerate(items):
    col, row = i % 3, i // 3
    card(s, 0.6 + col * 4.1, 1.8 + row * 2.45, 3.95, 2.25, t, b, c, 13)
phone(s, 'p-overview', 10.3, 4.55, 2.2) if False else None
rect(s, 8.8, 4.25, 3.95, 2.25, fill=PANEL); rect(s, 8.8, 4.25, 0.04, 2.25, ACC)
s.shapes.add_picture(LOGO, Inches(9.1), Inches(4.65), Inches(1.0), Inches(1.0))
text(s, 10.3, 4.8, 2.3, 0.9, ['LocalBoost', 'Observă · Diagnostichează', 'Acționează · Măsoară'], 12, MUTED, spacing=1.2)

# ---------------------------------------------------------------- 13 trust
s = new_slide('GIMI: pe scurt, partea tehnică. Funcționează offline ca PWA, există aplicație Android, AI-ul rulează pe server și are rezervă locală. Spunem deschis ce e simulat în demo.')
head(s, 'Funcții · Sub capotă', 'Încredere prin transparență', 13)
cells = [('PWA offline', 'Se instalează pe telefon și pe desktop, funcționează fără internet.'),
         ('Aplicație Android', 'Build Capacitor, același cod, ecran de pornire propriu.'),
         ('AI cu rezervă', 'Cheia rămâne pe server. Dacă AI-ul nu răspunde, intră regulile locale.'),
         ('5 limbi', 'Română, engleză, franceză, germană, maghiară. Temă luminoasă și întunecată.')]
for i, (t, b) in enumerate(cells):
    card(s, 0.6 + (i % 2) * 4.2, 1.8 + (i // 2) * 2.4, 4.05, 2.2, t, b, [ACC, ACC, GOLD, GOLD][i], 12)
phone(s, 'p-pulse', 9.8, 1.6, 4.9) if False else None
phone(s, 'p-overview', 9.8, 1.6, 4.9)
bar(s, 'Datele din demo sunt simulate și etichetate DEMO, INFERENȚĂ sau LIVE.') if False else None

# ---------------------------------------------------------------- 14 close
s = new_slide('TU: LocalBoost transformă ce se știe deja despre o firmă în pași care aduc clienți. Vă mulțumim, așteptăm întrebările.')
s.shapes.add_picture(LOGO, Inches(5.92), Inches(1.4), Inches(1.5), Inches(1.5))
text(s, 0, 3.2, 13.333, 0.9, 'Mulțumim!', 48, WHITE, True, align=PP_ALIGN.CENTER)
text(s, 0, 4.2, 13.333, 0.4, 'Din semnale împrăștiate, în pași care aduc clienți.', 20, GOLD, align=PP_ALIGN.CENTER)
rect(s, 6.33, 4.9, 0.67, 0.03, ACC)
text(s, 0, 5.2, 13.333, 0.4, 'Întrebări?', 16, MUTED, align=PP_ALIGN.CENTER)

prs.save(OUT); print('saved', OUT)
